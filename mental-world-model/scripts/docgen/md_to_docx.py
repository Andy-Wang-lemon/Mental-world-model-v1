#!/usr/bin/env python3
"""将 Markdown 转为 Word（段落、标题、表格、代码块）。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _set_run_font(run, *, bold: bool = False, mono: bool = False, size: int = 11) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    if mono:
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    else:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_rich_paragraph(doc: Document, text: str, *, style: str | None = None, mono: bool = False) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(6)

    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = p.add_run(text[pos : match.start()])
            _set_run_font(run, mono=mono)
        token = match.group(0)
        if token.startswith("**"):
            run = p.add_run(token[2:-2])
            _set_run_font(run, bold=True, mono=mono)
        else:
            run = p.add_run(token[1:-1])
            _set_run_font(run, mono=True)
        pos = match.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        _set_run_font(run, mono=mono)


def _parse_table_row(line: str) -> list[str] | None:
    line = line.strip()
    if not line.startswith("|") or line.count("|") < 2:
        return None
    cells = [c.strip() for c in line.strip("|").split("|")]
    return cells


def _is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) or c == "" for c in cells)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    headers = rows[0]
    body = rows[1:] if len(rows) > 1 else []
    table = doc.add_table(rows=1 + len(body), cols=len(headers))
    table.style = "Table Grid"
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        _set_run_font(run, bold=True)
    for row_idx, row in enumerate(body, start=1):
        for col, value in enumerate(row):
            cell = table.rows[row_idx].cells[col]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            _set_run_font(run)
    doc.add_paragraph()


def md_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)

    i = 0
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    while i < len(lines):
        line = lines[i]

        if in_code:
            if line.strip().startswith("```"):
                text = "\n".join(code_lines)
                if text.strip().startswith("flowchart"):
                    _add_rich_paragraph(
                        doc,
                        "【流程图】阶段 0 加载 → 阶段 1–3 提取（所缘 / 路由 / 想思）→ "
                        "阶段 4 评分 → 阶段 5 消解 → 阶段 6 种子 EMA 更新 → 下次交互复用 W_runtime。",
                        mono=True,
                    )
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Pt(12)
                    run = p.add_run(text)
                    _set_run_font(run, mono=True, size=10)
                    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                in_code = False
                code_lines = []
            else:
                code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("```"):
            in_code = True
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        row = _parse_table_row(line)
        if row is not None:
            if not _is_separator_row(row):
                table_rows.append(row)
            i += 1
            continue
        if table_rows:
            _add_table(doc, table_rows)
            table_rows = []

        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue

        _add_rich_paragraph(doc, line.strip())
        i += 1

    if table_rows:
        _add_table(doc, table_rows)

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def main() -> int:
    root = Path(__file__).resolve().parents[2]
    md_path = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "docs" / "execution_flow.md"
    docx_path = Path(sys.argv[2]) if len(sys.argv) > 2 else root / "definitions" / "raw" / "执行流程说明.docx"
    md_to_docx(md_path, docx_path)
    print(f"Wrote {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
