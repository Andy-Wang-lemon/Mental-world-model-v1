#!/usr/bin/env python3
"""生成《DynToM 与 RecToM 代跑说明》Word。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

_OUT = Path(__file__).resolve().parents[3] / "代跑命令说明.docx"


def _font(run, *, bold=False, size=11, mono=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if mono:
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    else:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def p(doc, text="", *, bold=False, size=11, space_after=8):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(0)
    if text:
        run = para.add_run(text)
        _font(run, bold=bold, size=size)
    return para


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def code_block(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = 1.15
    para.paragraph_format.space_after = Pt(10)
    para.paragraph_format.space_before = Pt(4)
    run = para.add_run(text)
    _font(run, mono=True, size=9, color=(0x22, 0x22, 0x22))


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        _font(run, bold=True, size=10)
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            _font(run, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def build(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(56)
    section.bottom_margin = Pt(56)
    section.left_margin = Pt(64)
    section.right_margin = Pt(64)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DynToM 与 RecToM 代跑说明")
    _font(run, bold=True, size=18)
    p(doc, "把压缩包解压后，按本文从上到下执行。不要改模型代码，不要删已有结果。")
    title.paragraph_format.space_after = Pt(12)

    heading(doc, "1. 这次要跑什么", 1)
    p(doc, "一共十个模型，两个数据集。都是在已经跑过的题目之外再加一批。")
    p(doc, "DynToM 再跑 18 条完整故事：trial50 到 trial67。每条故事里的题都要答，不是只答 18 道题。")
    p(doc, "RecToM 的推荐者、寻求者再各跑 180 题，也就是两个文件里下标 20 到 199。下标 0 到 19 已经有结果，不要重跑。")
    table(
        doc,
        ["模型", "接口", "温度", "DynToM 写到", "RecToM 写到"],
        [
            ["心智世界模型", "DeepSeek V4 Flash", "0", "reports/dyntom/dyntom mwm", "reports/rectom/mwm"],
            ["CoT", "DeepSeek V4 Flash", "0", "reports/dyntom/cot", "reports/rectom/cot"],
            ["SimToM", "DeepSeek V4 Flash", "0", "reports/dyntom/simtom", "reports/rectom/simtom"],
            ["TimeToM-style", "DeepSeek V4 Flash", "0", "reports/dyntom/timetom", "reports/rectom/timetom"],
            ["PICTURE", "DeepSeek V4 Flash", "0", "reports/dyntom/picture", "reports/rectom/picture"],
            ["Agentic-ToM", "DeepSeek V4 Flash", "0", "reports/dyntom/agentictom", "reports/rectom/agentictom"],
            ["DeepSeek 直接作答", "DeepSeek V4 Flash", "0", "reports/dyntom/dyntom_deepseek", "reports/rectom/rectom_deepseek"],
            ["Gemini", "gemini-3.8-flash", "0", "reports/dyntom/dyntom_gemini", "reports/rectom/rectom_gemini"],
            ["Grok", "grok-4.3", "0", "reports/dyntom/dyntom_grok", "reports/rectom/rectom_grok"],
            ["GPT", "gpt-6-luna", "1", "reports/dyntom/dyntom_gpt", "reports/rectom/rectom_gpt"],
        ],
    )
    p(doc, "前六个模型，再加上 DeepSeek 直接作答，共用同一个 DeepSeek 密钥。Gemini、Grok、GPT 各用自己的密钥。命令里的 APIKEY 要换成真实密钥，不要加引号，不要多空格。")

    heading(doc, "2. 解压后的目录", 1)
    p(doc, "压缩包解开后应能看到下面两层。命令都在 mental-world-model 里执行，数据在它旁边的公开数据集里。")
    code_block(
        doc,
        "心智世界模型（代码）/\n"
        "  mental-world-model/          ← 进入这一层再跑命令\n"
        "  公开数据集/\n"
        "    DynToM.json\n"
        "    RecTOM/3_pred_rec.json     ← 推荐者\n"
        "    RecTOM/4_pred_seeker.json  ← 寻求者",
    )
    p(doc, "请用 macOS、Linux，或 Windows 上的 Git Bash。不要用 cmd。下文的 export、for、seq 都是 bash 语法。")

    heading(doc, "3. 安装", 1)
    p(doc, "需要 Python 3.10 或更新版本。进入项目后安装依赖。若 python3 不存在，把下文所有 python3 换成 python。")
    code_block(
        doc,
        "cd mental-world-model\n"
        "python3 -m pip install -e .\n"
        "python3 -c \"import dotenv; print('ok')\"",
    )
    p(doc, "打印 ok 就可以开始。不要加 --mock。那是假模型，结果不能用。")

    heading(doc, "4. 每个终端先执行这一段", 1)
    p(doc, "下面先设 DeepSeek。换 Gemini、Grok、GPT 时，只改第 6 节里对应的四行 export，不必重做安装。TRIALS 和 args 在同一个终端里要保留。")
    code_block(
        doc,
        "cd mental-world-model\n"
        "\n"
        "export MWM_LLM_BASE_URL=\"https://api.deepseek.com\"\n"
        "export MWM_LLM_MODEL=\"deepseek-v4-flash\"\n"
        "export MWM_LLM_API_KEY=\"APIKEY\"\n"
        "export MWM_LLM_TEMPERATURE=\"0\"\n"
        "\n"
        "DATA=\"../公开数据集\"\n"
        "DYN=\"$DATA/DynToM.json\"\n"
        "REC=\"$DATA/RecTOM/3_pred_rec.json\"\n"
        "SEEKER=\"$DATA/RecTOM/4_pred_seeker.json\"\n"
        "TRIALS=(trial50 trial51 trial52 trial53 trial54 trial55 trial56 trial57 trial58 trial59 trial60 trial61 trial62 trial63 trial64 trial65 trial66 trial67)\n"
        "\n"
        "args=()\n"
        "for id in \"${TRIALS[@]}\"; do args+=(--trial \"$id\"); done",
    )
    p(doc, "DATA 必须是 ../公开数据集。五个推理框架如果不写 --data-dir，会到压缩包外面去找数据，然后报找不到文件。")

    heading(doc, "5. DeepSeek：心智世界模型和五个推理框架", 1)
    p(doc, "保持第 4 节的 DeepSeek 环境。可以一个终端跑完再开下一个，也可以分终端同时跑不同模型。同时跑时注意接口限速。", bold=False)
    heading(doc, "5.1 心智世界模型", 2)
    p(doc, "结果默认写到 reports 里已有的心智世界模型目录。不用再写 --output。跑完会自动改 summary.txt 里心智世界模型那一段，其它模型的段落不要手改。")
    code_block(
        doc,
        "python3 dyntom_adapter/__main__.py run --input \"$DYN\" \"${args[@]}\"\n"
        "\n"
        "for i in $(seq 20 199); do\n"
        "  python3 rectom_adapter/__main__.py run --input \"$REC\" --index \"$i\" || echo \"推荐者 $i 失败，继续\"\n"
        "done\n"
        "for i in $(seq 20 199); do\n"
        "  python3 rectom_adapter/__main__.py run --input \"$SEEKER\" --index \"$i\" || echo \"寻求者 $i 失败，继续\"\n"
        "done",
    )
    heading(doc, "5.2 五个推理框架", 2)
    p(doc, "这五个的温度由程序固定为 0，不用改环境变量。已有结果的题会自动跳过。RecToM 这条会先读入 0 到 199，已完成的 0 到 19 再跳过，所以一开始会停一阵，不是卡死。")
    code_block(
        doc,
        "LIST=$(IFS=,; echo \"${TRIALS[*]}\")\n"
        "for method in cot simtom timetom picture agentictom; do\n"
        "  python3 run_benchmark.py --dataset dyntom --method \"$method\" --trials \"$LIST\" --data-dir \"$DATA\"\n"
        "  python3 run_benchmark.py --dataset rectom --method \"$method\" --rectom-end 200 --data-dir \"$DATA\"\n"
        "done",
    )
    heading(doc, "5.3 DeepSeek 直接作答", 2)
    p(doc, "仍用第 4 节的 DeepSeek 密钥和温度 0。这一组必须自己写 --output，否则会报缺参数。")
    code_block(
        doc,
        "python3 dyntom_adapter/__main__.py baseline --input \"$DYN\" --output \"reports/dyntom/dyntom_deepseek\" \"${args[@]}\"\n"
        "for i in $(seq 20 199); do\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$REC\" --output \"reports/rectom/rectom_deepseek\" --index \"$i\" || echo \"推荐者 $i 失败，继续\"\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$SEEKER\" --output \"reports/rectom/rectom_deepseek\" --index \"$i\" || echo \"寻求者 $i 失败，继续\"\n"
        "done",
    )

    heading(doc, "6. Gemini、Grok、GPT", 1)
    p(doc, "换模型前，在同一个终端执行对应的四行 export。DATA、DYN、REC、SEEKER、TRIALS、args 不用重设。这三组都是直接作答，都必须带 --output。")
    heading(doc, "6.1 Gemini，温度 0", 2)
    code_block(
        doc,
        "export MWM_LLM_BASE_URL=\"https://generativelanguage.googleapis.com/v1beta/openai\"\n"
        "export MWM_LLM_MODEL=\"gemini-3.8-flash\"\n"
        "export MWM_LLM_API_KEY=\"APIKEY\"\n"
        "export MWM_LLM_TEMPERATURE=\"0\"\n"
        "\n"
        "python3 dyntom_adapter/__main__.py baseline --input \"$DYN\" --output \"reports/dyntom/dyntom_gemini\" \"${args[@]}\"\n"
        "for i in $(seq 20 199); do\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$REC\" --output \"reports/rectom/rectom_gemini\" --index \"$i\" || echo \"推荐者 $i 失败，继续\"\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$SEEKER\" --output \"reports/rectom/rectom_gemini\" --index \"$i\" || echo \"寻求者 $i 失败，继续\"\n"
        "done",
    )
    heading(doc, "6.2 Grok，温度 0", 2)
    code_block(
        doc,
        "export MWM_LLM_BASE_URL=\"https://api.x.ai/v1\"\n"
        "export MWM_LLM_MODEL=\"grok-4.3\"\n"
        "export MWM_LLM_API_KEY=\"APIKEY\"\n"
        "export MWM_LLM_TEMPERATURE=\"0\"\n"
        "\n"
        "python3 dyntom_adapter/__main__.py baseline --input \"$DYN\" --output \"reports/dyntom/dyntom_grok\" \"${args[@]}\"\n"
        "for i in $(seq 20 199); do\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$REC\" --output \"reports/rectom/rectom_grok\" --index \"$i\" || echo \"推荐者 $i 失败，继续\"\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$SEEKER\" --output \"reports/rectom/rectom_grok\" --index \"$i\" || echo \"寻求者 $i 失败，继续\"\n"
        "done",
    )
    heading(doc, "6.3 GPT，温度 1", 2)
    p(doc, "只有这一组温度是 1。程序可能打印当前温度不是 0 的警告，这是预期的，不要把它改回 0，也不要因此停掉。")
    code_block(
        doc,
        "export MWM_LLM_BASE_URL=\"https://api.openai.com/v1\"\n"
        "export MWM_LLM_MODEL=\"gpt-6-luna\"\n"
        "export MWM_LLM_API_KEY=\"APIKEY\"\n"
        "export MWM_LLM_TEMPERATURE=\"1\"\n"
        "\n"
        "python3 dyntom_adapter/__main__.py baseline --input \"$DYN\" --output \"reports/dyntom/dyntom_gpt\" \"${args[@]}\"\n"
        "for i in $(seq 20 199); do\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$REC\" --output \"reports/rectom/rectom_gpt\" --index \"$i\" || echo \"推荐者 $i 失败，继续\"\n"
        "  python3 rectom_adapter/__main__.py baseline --input \"$SEEKER\" --output \"reports/rectom/rectom_gpt\" --index \"$i\" || echo \"寻求者 $i 失败，继续\"\n"
        "done",
    )

    heading(doc, "7. 中断了怎么续", 1)
    p(doc, "不要换输出目录，不要删已经写出的文件。用原来的密钥和原来的命令再跑。")
    p(doc, "心智世界模型和四个直接作答的 DynToM：如果某条 trial 只写了一部分，在原来的命令末尾加上 --resume。例如心智世界模型是 python3 dyntom_adapter/__main__.py run --input \"$DYN\" \"${args[@]}\" --resume。四个直接作答同样把 --resume 加在最后。")
    p(doc, "五个推理框架：已完成的题会自动跳过，把原命令再执行一次即可。只有上次失败的题，才在命令末尾加 --resume。")
    p(doc, "RecToM 的循环也可以原样再跑。已经有 answer.json 的题会提示已存在并继续下一题，不会覆盖。")

    heading(doc, "8. 跑完交回什么", 1)
    p(doc, "把整个 mental-world-model/reports 目录交回，至少包含这次新写出的 trial50 到 trial67，以及 RecToM 下标 20 到 199。原来的 trial235、trial497 和下标 0 到 19 也要留着，不要删。")
    p(doc, "五个推理框架和四个直接作答不会自动改总表。心智世界模型会改 reports/dyntom/summary.txt 和 reports/rectom/summary.txt 里它自己的那一段。看到总表里其它模型的旧数字还在，是正常的。")

    heading(doc, "9. 不要做的事", 1)
    p(doc, "不要加 --mock。不要改 trial 名单，不要把 RecToM 从 0 开始重跑。不要把五个推理框架指到 Gemini、Grok 或 GPT，那五个只能用 deepseek-v4-flash。不要把 GPT 的温度改成 0。不要把结果写到新的文件夹名。不要把 API 密钥写进文件再打包发回。")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


if __name__ == "__main__":
    build(_OUT)
    print(_OUT)
